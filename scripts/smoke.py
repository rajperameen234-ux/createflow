r"""End-to-end generation check (no network, no cloud).

Run:  .venv\Scripts\python.exe -m scripts.smoke
"""

from __future__ import annotations

import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.content import build_outline  # noqa: E402
from app.design import build_design_system  # noqa: E402
from app.pipeline import generate  # noqa: E402
from app.schemas import GenerateRequest, StylePreset  # noqa: E402

BRIEF = GenerateRequest(
    project_name="Bean House Coffee",
    topic="Launching a modern coffee shop in the city centre, built for people who want a calm, "
          "high-quality cup and a place to sit for an hour.",
    audience="Young adults and local customers",
    goal="Fill 60 seats on weekday mornings within the first quarter of trading",
    tone="Premium, friendly, modern",
    presentation_length=8,
    presentation_type="marketing",
    visual_style="Editorial luxury",
    style_preset=StylePreset.EDITORIAL_LUXURY,
    primary_color="#3A2418",
    secondary_color="#E8DCC8",
    accent_color="#C9A227",
    brand_name="Bean House Coffee",
    tagline="Slow mornings, poured with intent.",
    extra_instructions="Keep the tone warm and unhurried.",
)


def main() -> int:
    started = time.perf_counter()
    design = build_design_system(BRIEF)
    outline = build_outline(BRIEF)
    print(f"design: {design.preset_label} accent={design.accent} on_primary={design.on_primary}")
    print(f"outline: {len(outline)} slides -> {[s.layout for s in outline]}")

    result = generate(BRIEF)
    elapsed = time.perf_counter() - started
    print(f"\ngenerated in {elapsed:.2f}s -> {result.workspace}")
    for asset in result.assets:
        print(f"  {asset.kind.value:12} {asset.filename:34} {asset.size/1024:8.1f} KB "
              f"{asset.width or '-'}x{asset.height or '-'}")

    from pptx import Presentation

    deck = next(a for a in result.assets if a.kind.value == "presentation")
    prs = Presentation(str(deck.path))
    print(f"\nPPTX: {len(prs.slides)} slides, {prs.slide_width.inches:.2f}x{prs.slide_height.inches:.2f} in")
    total_shapes = 0
    charts = 0
    notes = 0
    for i, slide in enumerate(prs.slides, 1):
        shapes = len(slide.shapes)
        total_shapes += shapes
        has_chart = any(s.has_chart for s in slide.shapes)
        charts += int(has_chart)
        has_notes = slide.has_notes_slide and bool(slide.notes_slide.notes_text_frame.text.strip())
        notes += int(has_notes)
        titles = [s.text_frame.text for s in slide.shapes
                  if s.has_text_frame and s.text_frame.text.strip()]
        head = titles[1][:52] if len(titles) > 1 else (titles[0][:52] if titles else "")
        print(f"  s{i:02d} shapes={shapes:3} chart={'Y' if has_chart else 'n'} "
              f"notes={'Y' if has_notes else 'n'} | {head}")
    print(f"\ntotal shapes={total_shapes} charts={charts} slides_with_notes={notes}")

    from PIL import Image

    for kind in ("poster", "social"):
        asset = next(a for a in result.assets if a.kind.value == kind)
        with Image.open(asset.path) as img:
            img.verify()
        with Image.open(asset.path) as img:
            print(f"{kind}: opens OK, mode={img.mode}, size={img.size}, format={img.format}")

    svg = next(a for a in result.assets if a.kind.value == "svg")
    print(f"svg: {len(svg.path.read_text(encoding='utf-8'))} bytes, contains <svg>: "
          f"{'<svg' in svg.path.read_text(encoding='utf-8')}")

    keep = Path(__file__).resolve().parent.parent / ".smoke-output"
    keep.mkdir(exist_ok=True)
    for asset in result.assets:
        (keep / asset.filename).write_bytes(asset.path.read_bytes())
    print(f"\ncopied samples to {keep}")

    result.cleanup()
    print(f"temp workspace removed: {not result.workspace.exists()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())