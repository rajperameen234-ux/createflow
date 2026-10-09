"""PowerPoint generation with python-pptx.

Everything drawn here is a native, editable PowerPoint object: text frames,
autoshapes, connectors and a real embedded chart part. No flattened images.
"""

from __future__ import annotations

import math
from pathlib import Path
from typing import Optional

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from .content import Slide, build_outline, headline_from, short, title_case
from .design import (
    DesignSystem,
    build_design_system,
    contrast_ratio,
    ensure_contrast,
    mix,
    shade,
)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.85)
CONTENT_TOP = Inches(1.95)
CONTENT_BOTTOM = Inches(6.85)


def rgb(color: str) -> RGBColor:
    value = color.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


class DeckBuilder:
    def __init__(self, brief) -> None:
        self.brief = brief
        self.design: DesignSystem = build_design_system(brief)
        self.prs = Presentation()
        self.prs.slide_width = SLIDE_W
        self.prs.slide_height = SLIDE_H
        self.blank = self.prs.slide_layouts[6]
        self.index = 0
        self.total = 0
        self.current_bg = self.design.paper

    # -- primitives ------------------------------------------------------
    def new_slide(self, bg: Optional[str] = None) -> "SlideBox":
        self.index += 1
        base = bg or self.design.paper
        self.current_bg = base
        slide = self.prs.slides.add_slide(self.blank)
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb(base)
        return SlideBox(self, slide)

    def rect(self, slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE,
             line_width: float = 1.0, adj: Optional[float] = None, shadow: bool = False):
        shp = slide.shapes.add_shape(shape, int(x), int(y), int(w), int(h))
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(fill)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = rgb(line)
            shp.line.width = Pt(line_width)
        else:
            shp.line.fill.background()
        if not shadow:
            self._no_shadow(shp)
        if adj is not None:
            try:
                shp.adjustments[0] = adj
            except (IndexError, ValueError):
                pass
        shp.text_frame.word_wrap = True
        return shp

    @staticmethod
    def _no_shadow(shp) -> None:
        try:
            shp.shadow.inherit = False
        except (AttributeError, NotImplementedError):
            pass

    def textbox(self, slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
        box = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Emu(0)
        tf.margin_top = tf.margin_bottom = Emu(0)
        tf.vertical_anchor = anchor
        return box, tf

    def write(self, tf, text, *, size, color, font=None, bold=False, italic=False,
              align=PP_ALIGN.LEFT, space_after=0, line_spacing=None, caps=False,
              tracking: Optional[float] = None, first=False):
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        para.alignment = align
        if space_after:
            para.space_after = Pt(space_after)
        if line_spacing:
            para.line_spacing = line_spacing
        run = para.add_run()
        run.text = text.upper() if caps else text
        f = run.font
        f.size = Pt(size)
        f.bold = bold
        f.italic = italic
        f.name = font or self.design.body_font
        f.color.rgb = rgb(color)
        if tracking:
            # python-pptx has no tracking API; set it on the run properties directly.
            try:
                rPr = run._r.get_or_add_rPr()
                rPr.set("spc", str(int(tracking * 100)))
            except (AttributeError, ValueError):
                pass
        return para

    # -- shared furniture ------------------------------------------------
    def frame(self, slide, kicker: str, title: str, *, bg: Optional[str] = None,
              title_color: Optional[str] = None) -> None:
        d = self.design
        base = bg or d.paper
        self.index = self.index
        label_color = d.accent if contrast_ratio(d.accent, base) >= 3 else ensure_contrast(d.accent, base, 3)

        if kicker:
            _, tf = self.textbox(slide, MARGIN, Inches(0.72), SLIDE_W - MARGIN * 2, Inches(0.28))
            self.write(tf, kicker, size=d.type_scale["micro"], color=label_color,
                       bold=True, caps=True, tracking=1.6, first=True)

        title_color = title_color or d.ink
        _, tf = self.textbox(slide, MARGIN, Inches(1.05), SLIDE_W - MARGIN * 2, Inches(0.85))
        self.write(tf, title, size=d.type_scale["section"], color=title_color,
                   font=d.heading_font, bold=True, line_spacing=1.0,
                   tracking=d.heading_tracking, first=True)

        rule_w = Inches(1.5)
        self.rect(slide, MARGIN, Inches(1.74), rule_w, Pt(2.5 * d.rule_weight),
                  fill=d.accent, shape=MSO_SHAPE.RECTANGLE)

        self.pattern_band(slide)

    def pattern_band(self, slide) -> None:
        d = self.design
        band_y = SLIDE_H - Inches(0.52)
        base_hex = self.current_bg
        hairline = mix(base_hex, d.ink, 0.18)
        self.rect(slide, MARGIN, band_y, SLIDE_W - MARGIN * 2, Emu(9525), fill=hairline)
        _, tf = self.textbox(slide, MARGIN, band_y + Inches(0.1), Inches(6), Inches(0.24))
        self.write(tf, f"{d.brand_name} · {d.preset_label}", size=8.5,
                   color=mix(base_hex, d.ink, 0.55), bold=True, caps=True, tracking=1.2, first=True)
        num = f"{self.index:02d} / {self.total:02d}"
        _, tf2 = self.textbox(slide, SLIDE_W - MARGIN - Inches(2), band_y + Inches(0.1),
                              Inches(2), Inches(0.24))
        self.write(tf2, num, size=8.5, color=mix(base_hex, d.ink, 0.55),
                   align=PP_ALIGN.RIGHT, bold=True, tracking=1.2, first=True)

    def notes(self, slide, text: str) -> None:
        if not text:
            return
        try:
            slide.notes_slide.notes_text_frame.text = text
        except (AttributeError, NotImplementedError):
            pass

    def logo_mark(self, slide, x, y, size, *, on_dark=False) -> None:
        d = self.design
        base = d.ink if not on_dark else d.paper
        self.rect(slide, x, y, size, size, fill=None, line=d.accent, shape=MSO_SHAPE.OVAL,
                  line_width=1.5)
        _, tf = self.textbox(slide, x, y + size * 0.26, size, size * 0.5,
                             anchor=MSO_ANCHOR.MIDDLE)
        self.write(tf, d.brand_name[:2].upper(), size=12, color=d.accent,
                   font=d.heading_font, bold=True, align=PP_ALIGN.CENTER, first=True)
        del base


class SlideBox:
    """Thin proxy so layout functions read as ``slide.shape``."""

    def __init__(self, builder: DeckBuilder, native) -> None:
        self.b = builder
        self.shapes = native.shapes
        self.native = native


def _card(builder: DeckBuilder, slide, x, y, w, h, *, fill: Optional[str] = None,
          line: Optional[str] = None, radius: Optional[int] = None) -> None:
    d = builder.design
    fill = fill or d.surface
    line = line or mix(d.paper, d.ink, 0.12)
    builder.rect(slide, x, y, w, h, fill=fill, line=line,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                 adj=(radius if radius is not None else d.corner_radius) / 100.0)


# --------------------------------------------------------------------------
# layouts
# --------------------------------------------------------------------------
def layout_hero(b: DeckBuilder, s: Slide, total: int) -> SlideBox:
    d = b.design
    bg = d.primary if not d.dark_canvas else shade(d.paper, -0.02)
    slide = b.new_slide(bg)

    # motif: concentric arcs bleeding off the right edge. Named "bleed-" so the
    # deck validator can tell intentional full-bleed decoration from mistakes.
    cx = SLIDE_W - Inches(2.1)
    cy = SLIDE_H * 0.52
    for i, ratio in enumerate((0.55, 0.8, 1.05, 1.3)):
        size = Inches(6.2 * ratio)
        arc = b.rect(slide, cx - size / 2, cy - size / 2, size, size, fill=None,
                     line=mix(bg, d.accent, 0.55 - i * 0.08), shape=MSO_SHAPE.OVAL,
                     line_width=1.1)
        arc.name = f"bleed-arc-{i + 1}"
    b.rect(slide, cx - Inches(1.05), cy - Inches(1.05), Inches(2.1), Inches(2.1),
           fill=mix(bg, d.accent, 0.22), shape=MSO_SHAPE.OVAL)

    left_w = Inches(7.6)
    x = MARGIN
    if s.kicker:
        _, tf = b.textbox(slide, x, Inches(1.5), left_w, Inches(0.3))
        b.write(tf, s.kicker, size=11, color=mix(bg, d.accent, 0.85), bold=True,
                caps=True, tracking=2.0, first=True)

    _, tf = b.textbox(slide, x, Inches(2.0), left_w, Inches(2.4))
    hero_color = ensure_contrast(d.paper if d.dark_canvas else "#FFFFFF", bg, 8.0)
    b.write(tf, s.title, size=d.type_scale["hero"], color=hero_color, font=d.heading_font,
            bold=True, line_spacing=0.95, tracking=d.heading_tracking, first=True)

    b.rect(slide, x, Inches(4.5), Inches(1.7), Pt(3 * d.rule_weight), fill=d.accent)

    if s.body:
        _, tf = b.textbox(slide, x, Inches(4.85), Inches(6.4), Inches(1.1))
        b.write(tf, s.body, size=d.type_scale["subtitle"], color=mix(bg, hero_color, 0.35),
                line_spacing=1.25, first=True)

    b.logo_mark(slide, x, Inches(6.1), Inches(0.62), on_dark=True)
    _, tf = b.textbox(slide, x + Inches(0.85), Inches(6.14), Inches(5), Inches(0.6))
    b.write(tf, d.brand_name, size=13, color=hero_color, font=d.heading_font, bold=True, first=True)
    b.write(tf, title_case(short(d.tagline, 60)), size=10.5,
            color=mix(bg, hero_color, 0.45), caps=True, tracking=1.4)

    b.notes(slide.native, s.notes)
    return slide


def layout_two_column(b: DeckBuilder, s: Slide, total: int) -> SlideBox:
    d = b.design
    slide = b.new_slide()
    b.frame(slide, s.kicker, s.title)

    gap = Inches(0.6)
    col_w = (SLIDE_W - MARGIN * 2 - gap) / 2
    top = CONTENT_TOP
    body_h = Inches(2.5) if s.bullets else Inches(4.3)

    # left: narrative paragraph on a tinted panel
    panel_h = body_h + Inches(0.2)
    _card(b, slide, MARGIN, top, col_w, panel_h, fill=mix(d.paper, d.secondary, 0.55))
    b.rect(slide, MARGIN, top, Pt(3), panel_h, fill=d.accent)
    _, tf = b.textbox(slide, MARGIN + Inches(0.42), top + Inches(0.34),
                      col_w - Inches(0.72), panel_h - Inches(0.6))
    b.write(tf, s.body or s.title, size=d.type_scale["body"] + 1, color=d.ink,
            line_spacing=1.42, first=True)

    # right: bullet cards
    rx = MARGIN + col_w + gap
    items = s.bullets[:4] or []
    card_h = (panel_h - Inches(0.2) * (len(items) - 1)) / max(1, len(items))
    for i, item in enumerate(items):
        y = top + (card_h + Inches(0.2)) * i
        _card(b, slide, rx, y, col_w, card_h, fill=d.surface)
        b.rect(slide, rx + Inches(0.3), y + Inches(0.3), Inches(0.34), Inches(0.34),
               fill=d.accent, shape=MSO_SHAPE.OVAL)
        _, tf = b.textbox(slide, rx + Inches(0.3), y + Inches(0.36), Inches(0.34), Inches(0.3))
        b.write(tf, str(i + 1), size=11, color=ensure_contrast("#FFFFFF", d.accent, 4.0),
                font=d.heading_font, bold=True, align=PP_ALIGN.CENTER, first=True)
        _, tf = b.textbox(slide, rx + Inches(0.8), y + Inches(0.28),
                          col_w - Inches(1.1), card_h - Inches(0.5))
        b.write(tf, item.head, size=d.type_scale["body"], color=d.ink,
                font=d.heading_font, bold=True, first=True)
        b.write(tf, item.body, size=d.type_scale["small"] + 1, color=d.muted, line_spacing=1.25,
                space_after=0)

    b.pattern_band(slide)
    if s.caption:
        _, tf = b.textbox(slide, MARGIN, CONTENT_BOTTOM + Inches(0.06),
                          SLIDE_W - MARGIN * 2, Inches(0.3))
        b.write(tf, s.caption, size=9.5, color=d.muted, italic=True, first=True)
    b.notes(slide.native, s.notes)
    return slide


def layout_numbered_process(b: DeckBuilder, s: Slide, total: int) -> SlideBox:
    d = b.design
    slide = b.new_slide()
    b.frame(slide, s.kicker, s.title)

    steps = s.steps[:5]
    n = len(steps) or 1
    gap = Inches(0.28)
    col_w = (SLIDE_W - MARGIN * 2 - gap * (n - 1)) / n
    top = CONTENT_TOP + Inches(0.55)
    card_h = Inches(2.5)

    rail_y = top - Inches(0.34)
    b.rect(slide, MARGIN + Inches(0.2), rail_y, SLIDE_W - MARGIN * 2 - Inches(0.4),
           Emu(9525), fill=mix(d.paper, d.ink, 0.15))

    for i, (head, body) in enumerate(steps):
        x = MARGIN + (col_w + gap) * i
        b.rect(slide, x + col_w / 2 - Inches(0.16), rail_y - Inches(0.15), Inches(0.32), Inches(0.32),
               fill=d.accent, shape=MSO_SHAPE.OVAL)
        _card(b, slide, x, top, col_w, card_h, fill=d.surface)
        b.rect(slide, x, top, col_w, Inches(0.16), fill=d.accent if i == 0 else d.primary,
               shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.25)
        _, tf = b.textbox(slide, x + Inches(0.26), top + Inches(0.42), col_w - Inches(0.52), Inches(0.4))
        b.write(tf, f"{i + 1:02d}", size=26, color=mix(d.paper, d.accent, 0.75),
                font=d.heading_font, bold=True, first=True)
        _, tf = b.textbox(slide, x + Inches(0.26), top + Inches(0.95), col_w - Inches(0.52), Inches(0.4))
        b.write(tf, title_case(head), size=d.type_scale["body"] + 1, color=d.ink,
                font=d.heading_font, bold=True, first=True)
        _, tf = b.textbox(slide, x + Inches(0.26), top + Inches(1.38), col_w - Inches(0.52),
                          card_h - Inches(1.6))
        b.write(tf, body, size=d.type_scale["small"], color=d.muted, line_spacing=1.3, first=True)

    if s.body:
        _, tf = b.textbox(slide, MARGIN, CONTENT_TOP - Inches(0.05), Inches(9.5), Inches(0.4))
        b.write(tf, s.body, size=d.type_scale["body"], color=d.muted, first=True)

    b.pattern_band(slide)
    if s.caption:
        _, tf = b.textbox(slide, MARGIN, CONTENT_BOTTOM + Inches(0.06),
                          SLIDE_W - MARGIN * 2, Inches(0.3))
        b.write(tf, s.caption, size=9.5, color=d.muted, italic=True, first=True)
    b.notes(slide.native, s.notes)
    return slide


def layout_timeline(b: DeckBuilder, s: Slide, total: int) -> SlideBox:
    d = b.design
    slide = b.new_slide()
    b.frame(slide, s.kicker, s.title)

    items = s.timeline[:6]
    n = len(items) or 1
    rail_y = CONTENT_TOP + Inches(1.0)
    x0 = MARGIN + Inches(0.4)
    x1 = SLIDE_W - MARGIN - Inches(0.4)
    b.rect(slide, x0, rail_y, x1 - x0, Pt(2), fill=mix(d.paper, d.ink, 0.18))

    step = (x1 - x0) / max(1, n - 1) if n > 1 else 0
    for i, (label, body) in enumerate(items):
        cx = int(x0 + step * i)
        b.rect(slide, cx - Inches(0.19), rail_y - Inches(0.17), Inches(0.38), Inches(0.38),
               fill=d.accent, shape=MSO_SHAPE.OVAL)
        b.rect(slide, cx - Inches(0.08), rail_y - Inches(0.06), Inches(0.16), Inches(0.16),
               fill=ensure_contrast("#FFFFFF", d.accent, 3.0), shape=MSO_SHAPE.OVAL)

        above = i % 2 == 0
        card_w = Inches(2.75)
        card_h = Inches(1.35)
        card_x = min(max(cx - card_w / 2, MARGIN), SLIDE_W - MARGIN - card_w)
        card_y = rail_y - Inches(0.55) - card_h if above else rail_y + Inches(0.55)
        _card(b, slide, card_x, card_y, card_w, card_h, fill=d.surface)

        conn_y = card_y + card_h if above else rail_y + Inches(0.2)
        line = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, cx, rail_y, cx, int(conn_y)
        )
        line.line.color.rgb = rgb(mix(d.paper, d.ink, 0.3))
        line.line.width = Pt(0.75)

        _, tf = b.textbox(slide, card_x + Inches(0.24), card_y + Inches(0.2),
                          card_w - Inches(0.48), Inches(0.28))
        b.write(tf, label, size=10, color=d.accent if contrast_ratio(d.accent, d.paper) >= 3 else d.ink,
                bold=True, caps=True, tracking=1.4, first=True)
        _, tf = b.textbox(slide, card_x + Inches(0.24), card_y + Inches(0.52),
                          card_w - Inches(0.48), card_h - Inches(0.7))
        b.write(tf, body, size=d.type_scale["small"], color=d.ink, line_spacing=1.28, first=True)

    b.pattern_band(slide)
    if s.caption:
        _, tf = b.textbox(slide, MARGIN, CONTENT_BOTTOM + Inches(0.06),
                          SLIDE_W - MARGIN * 2, Inches(0.3))
        b.write(tf, s.caption, size=9.5, color=d.muted, italic=True, first=True)
    b.notes(slide.native, s.notes)
    return slide


def _metric_card(b: DeckBuilder, slide, x, y, w, h, metric, index: int) -> None:
    d = b.design
    _card(b, slide, x, y, w, h, fill=d.surface)
    b.rect(slide, x, y, w, Pt(3), fill=d.accent if index % 2 == 0 else d.primary,
           shape=MSO_SHAPE.RECTANGLE)
    _, tf = b.textbox(slide, x + Inches(0.3), y + Inches(0.42), w - Inches(0.6), Inches(0.9))
    b.write(tf, metric.value, size=d.type_scale["metric"], color=d.accent
            if contrast_ratio(d.accent, d.surface) >= 3 else d.ink,
            font=d.heading_font, bold=True, first=True)
    _, tf = b.textbox(slide, x + Inches(0.3), y + Inches(1.28), w - Inches(0.6), Inches(0.4))
    b.write(tf, title_case(metric.label), size=d.type_scale["body"], color=d.ink,
            font=d.heading_font, bold=True, first=True)
    if metric.note:
        _, tf = b.textbox(slide, x + Inches(0.3), y + Inches(1.68), w - Inches(0.6), h - Inches(1.85))
        b.write(tf, metric.note, size=d.type_scale["small"] - 1, color=d.muted, line_spacing=1.25,
                first=True)


def layout_metrics(b: DeckBuilder, s: Slide, total: int) -> SlideBox:
    d = b.design
    slide = b.new_slide()
    b.frame(slide, s.kicker, s.title)

    if s.body:
        _, tf = b.textbox(slide, MARGIN, CONTENT_TOP - Inches(0.02), Inches(10.4), Inches(0.8))
        b.write(tf, s.body, size=d.type_scale["body"], color=d.muted, line_spacing=1.35, first=True)
        content_top = CONTENT_TOP + Inches(0.92)
    else:
        content_top = CONTENT_TOP

    metrics = s.metrics[:4]
    gap = Inches(0.4)
    card_w = (SLIDE_W - MARGIN * 2 - gap * 3) / 4
    card_h = Inches(1.9)
    for i, metric in enumerate(metrics):
        _metric_card(b, slide, MARGIN + (card_w + gap) * i, content_top, card_w, card_h, metric, i)

    if s.series:
        used = content_top + (card_h + Inches(0.45) if metrics else 0)
        chart_y = used
        chart_h = SLIDE_H - chart_y - Inches(0.95)
        if chart_h > Inches(1.4):
            b.add_chart(slide, s, MARGIN, chart_y, SLIDE_W - MARGIN * 2, chart_h)
        elif metrics:
            # Not enough vertical room: drop the metric strip and chart full height.
            for shape in list(slide.shapes):
                if shape.top is not None and content_top <= shape.top < used:
                    shape._element.getparent().remove(shape._element)
            b.add_chart(slide, s, MARGIN, content_top, SLIDE_W - MARGIN * 2,
                        SLIDE_H - content_top - Inches(0.95))

    b.pattern_band(slide)
    if s.caption:
        _, tf = b.textbox(slide, MARGIN, SLIDE_H - Inches(0.82), SLIDE_W - MARGIN * 2, Inches(0.3))
        b.write(tf, s.caption, size=9.5,
                color=d.accent if "Illustrative" in s.caption else d.muted,
                italic="Illustrative" not in s.caption, bold="Illustrative" in s.caption, first=True)
    b.notes(slide.native, s.notes)
    return slide


def _add_chart_impl(b: DeckBuilder, slide, s: Slide, x, y, w, h) -> None:
    d = b.design
    series = s.series[:3]
    if not series or not series[0].points:
        return
    categories = [p[0] for p in series[0].points]
    data = CategoryChartData()
    data.categories = categories
    for ser in series:
        data.add_series(ser.name, tuple(float(v) for _, v in ser.points))

    kinds = {
        "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "bar": XL_CHART_TYPE.BAR_CLUSTERED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "doughnut": XL_CHART_TYPE.DOUGHNUT,
    }
    gfx = slide.shapes.add_chart(kinds[s.chart_kind], int(x), int(y), int(w), int(h), data)
    chart = gfx.chart

    chart.has_title = True
    chart.chart_title.text_frame.text = s.chart_title or "Illustrative model"
    tr = chart.chart_title.text_frame.paragraphs[0].runs[0]
    tr.font.size = Pt(12)
    tr.font.bold = True
    tr.font.name = d.heading_font
    tr.font.color.rgb = rgb(d.ink)

    chart.has_legend = len(series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(10)
        chart.legend.font.name = d.body_font
        chart.legend.font.color.rgb = rgb(d.muted)

    try:
        chart.font.size = Pt(10)
        chart.font.name = d.body_font
        chart.font.color.rgb = rgb(d.muted)
    except (AttributeError, NotImplementedError):
        pass

    palette = d.series(len(series))
    for i, point in enumerate(chart.series):
        try:
            if s.chart_kind == "line":
                point.format.line.color.rgb = rgb(palette[i])
                point.format.line.width = Pt(2.25)
                point.smooth = False
            else:
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = rgb(palette[i])
                point.format.line.fill.background()
        except (AttributeError, NotImplementedError, ValueError):
            pass

    if s.chart_kind in ("column", "bar"):
        plot = chart.plots[0]
        plot.gap_width = 60
        plot.overlap = -10 if len(series) > 1 else 0
        try:
            plot.has_data_labels = True
            labels = plot.data_labels
            labels.font.size = Pt(9)
            labels.font.name = d.body_font
            labels.font.color.rgb = rgb(d.muted)
            labels.position = XL_LABEL_POSITION.OUTSIDE_END
            labels.number_format = '0'
            labels.number_format_is_linked = False
        except (AttributeError, NotImplementedError, ValueError):
            pass

    try:
        va = chart.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = rgb(mix(d.paper, d.ink, 0.14))
        va.major_gridlines.format.line.width = Pt(0.75)
        va.format.line.color.rgb = rgb(mix(d.paper, d.ink, 0.22))
        va.tick_labels.font.size = Pt(9)
        ca = chart.category_axis
        ca.format.line.color.rgb = rgb(mix(d.paper, d.ink, 0.22))
        ca.tick_labels.font.size = Pt(9)
        ca.tick_labels.font.bold = True
    except (AttributeError, NotImplementedError, ValueError):
        pass


DeckBuilder.add_chart = _add_chart_impl  # type: ignore[attr-defined]


def layout_comparison(b: DeckBuilder, s: Slide, total: int) -> SlideBox:
    d = b.design
    slide = b.new_slide()
    b.frame(slide, s.kicker, s.title)

    before, after, impact = s.compare
    if not before:
        before, after = ["Current state"], ["Proposed state"]

    rows = max(len(before), len(after), 1)
    top = CONTENT_TOP + Inches(0.42)
    table_h = SLIDE_H - top - Inches(1.15)
    gap = Inches(0.35)
    label_w = Inches(2.5)
    col_w = (SLIDE_W - MARGIN * 2 - label_w - gap * 2) / 2
    header_h = Inches(0.52)
    row_h = (table_h - header_h) / rows

    _, tf = b.textbox(slide, MARGIN, top - Inches(0.32), label_w, Inches(0.28))
    b.write(tf, "Dimension", size=9.5, color=d.muted, bold=True, caps=True, tracking=1.4, first=True)

    for col, (title, x, fill) in enumerate(
        (
            (s.compare_titles[0], MARGIN + label_w + gap, mix(d.paper, d.ink, 0.06)),
            (s.compare_titles[1], MARGIN + label_w + gap * 2 + col_w, d.primary),
        )
    ):
        b.rect(slide, x, top, col_w, header_h, fill=fill,
               shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.18)
        _, tf = b.textbox(slide, x + Inches(0.22), top + Inches(0.13), col_w - Inches(0.44),
                          Inches(0.3))
        b.write(tf, title_case(title), size=11, color=ensure_contrast(d.paper, fill, 4.0),
                bold=True, caps=True, tracking=1.2, first=True)
        del col

    for i in range(rows):
        y = top + header_h + row_h * i
        if i % 2 == 0:
            b.rect(slide, MARGIN, y, SLIDE_W - MARGIN * 2, row_h,
                   fill=mix(d.paper, d.surface, 0.9))

        label = impact[i] if i < len(impact) else ""
        _, tf = b.textbox(slide, MARGIN + Inches(0.1), y + Inches(0.1), label_w - Inches(0.2),
                          row_h - Inches(0.18))
        b.write(tf, title_case(label), size=10.5, color=d.ink, bold=True, line_spacing=1.15,
                first=True)

        left = before[i] if i < len(before) else ""
        right = after[i] if i < len(after) else ""
        lx = MARGIN + label_w + gap
        rx = lx + col_w + gap

        b.rect(slide, lx, y + Inches(0.16), Inches(0.22), Inches(0.22), fill=None,
               line=mix(d.ink, d.paper, 0.45), shape=MSO_SHAPE.OVAL, line_width=1.0)
        _, tf = b.textbox(slide, lx + Inches(0.4), y + Inches(0.14), col_w - Inches(0.55),
                          row_h - Inches(0.24))
        b.write(tf, left, size=10.5, color=d.muted, line_spacing=1.18, first=True)

        b.rect(slide, rx, y + Inches(0.16), Inches(0.22), Inches(0.22), fill=d.accent,
               shape=MSO_SHAPE.OVAL)
        _, tf = b.textbox(slide, rx + Inches(0.4), y + Inches(0.14), col_w - Inches(0.55),
                          row_h - Inches(0.24))
        b.write(tf, right, size=10.5, color=d.ink, line_spacing=1.18, bold=True, first=True)

    b.pattern_band(slide)
    if s.caption:
        _, tf = b.textbox(slide, MARGIN, CONTENT_BOTTOM + Inches(0.06),
                          SLIDE_W - MARGIN * 2, Inches(0.3))
        b.write(tf, s.caption, size=9.5, color=d.muted, italic=True, first=True)
    b.notes(slide.native, s.notes)
    return slide


def layout_framework(b: DeckBuilder, s: Slide, total: int) -> SlideBox:
    d = b.design
    slide = b.new_slide()
    b.frame(slide, s.kicker, s.title)

    items = s.quadrant[:4]
    n = len(items) or 1
    rows = 2
    gap = Inches(0.4)
    card_w = (SLIDE_W - MARGIN * 2 - gap) / 2
    top = CONTENT_TOP + Inches(0.5)
    card_h = (SLIDE_H - top - Inches(1.3) - gap) / rows

    for i, (num, head, body) in enumerate(items):
        r, c = divmod(i, 2)
        x = MARGIN + (card_w + gap) * c
        y = top + (card_h + gap) * r
        _card(b, slide, x, y, card_w, card_h, fill=d.surface)
        b.rect(slide, x, y, Inches(0.18), card_h, fill=d.accent if i % 2 == 0 else d.primary,
               shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.4)
        _, tf = b.textbox(slide, x + Inches(0.44), y + Inches(0.3), Inches(1.2), Inches(0.5))
        b.write(tf, num, size=30, color=mix(d.paper, d.accent, 0.6), font=d.heading_font,
                bold=True, first=True)
        _, tf = b.textbox(slide, x + Inches(1.5), y + Inches(0.34), card_w - Inches(1.9), Inches(0.4))
        b.write(tf, title_case(head), size=d.type_scale["subtitle"], color=d.ink,
                font=d.heading_font, bold=True, first=True)
        _, tf = b.textbox(slide, x + Inches(1.5), y + Inches(0.86), card_w - Inches(1.9),
                          card_h - Inches(1.1))
        b.write(tf, body, size=d.type_scale["small"] + 1, color=d.muted, line_spacing=1.3, first=True)

    if s.body:
        _, tf = b.textbox(slide, MARGIN, CONTENT_TOP - Inches(0.02), Inches(10.4), Inches(0.45))
        b.write(tf, s.body, size=d.type_scale["body"], color=d.muted, first=True)

    b.pattern_band(slide)
    if s.caption:
        _, tf = b.textbox(slide, MARGIN, CONTENT_BOTTOM + Inches(0.06),
                          SLIDE_W - MARGIN * 2, Inches(0.3))
        b.write(tf, s.caption, size=9.5, color=d.muted, italic=True, first=True)
    b.notes(slide.native, s.notes)
    return slide


def layout_quote_led(b: DeckBuilder, s: Slide, total: int) -> SlideBox:
    d = b.design
    slide = b.new_slide()
    b.frame(slide, s.kicker, s.title)

    panel_x = MARGIN
    panel_w = Inches(8.2)
    panel_y = CONTENT_TOP + Inches(0.3)
    panel_h = SLIDE_H - panel_y - Inches(1.25)
    _card(b, slide, panel_x, panel_y, panel_w, panel_h, fill=d.primary)
    b.rect(slide, panel_x + Inches(0.5), panel_y + Inches(0.45), Inches(0.9),
           Pt(3 * d.rule_weight), fill=d.accent)

    _, tf = b.textbox(slide, panel_x + Inches(0.5), panel_y + Inches(0.85),
                      panel_w - Inches(1.0), panel_h - Inches(1.9))
    on_dark = ensure_contrast(d.paper, d.primary, 8.0)
    b.write(tf, s.body or s.title, size=d.type_scale["section"], color=on_dark,
            font=d.heading_font, bold=True, italic=True, line_spacing=1.24, first=True)

    if s.media_label:
        _, tf = b.textbox(slide, panel_x + Inches(0.5), panel_y + panel_h - Inches(0.85),
                          panel_w - Inches(1.0), Inches(0.5))
        b.write(tf, "— " + title_case(s.media_label), size=11,
                color=mix(d.primary, on_dark, 0.55), caps=True, tracking=1.4, first=True)

    # decorative mark column
    rx = panel_x + panel_w + Inches(0.55)
    rw = SLIDE_W - MARGIN - rx
    for i, ratio in enumerate((0.9, 0.65, 0.4)):
        size = int(rw * ratio)
        b.rect(slide, rx + (rw - size) / 2, panel_y + (panel_h - size) / 2, size, size,
               fill=None, line=mix(d.paper, d.accent, 0.6 - i * 0.12), shape=MSO_SHAPE.OVAL,
               line_width=1.2)

    b.pattern_band(slide)
    if s.caption:
        _, tf = b.textbox(slide, MARGIN, CONTENT_BOTTOM + Inches(0.06),
                          SLIDE_W - MARGIN * 2, Inches(0.3))
        b.write(tf, s.caption, size=9.5, color=d.muted, italic=True, first=True)
    b.notes(slide.native, s.notes)
    return slide


def _media_panel(b: DeckBuilder, slide, x, y, w, h, label: str) -> None:
    """Vector 'image-led' panel: motif shapes instead of a photo."""
    d = b.design
    b.rect(slide, x, y, w, h, fill=d.primary, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
           adj=d.corner_radius / 100.0)
    inset = Inches(0.45)
    inner_x, inner_y = x + inset, y + inset
    inner_w, inner_h = w - inset * 2, h - inset * 2
    ring = min(inner_w, inner_h)

    b.rect(slide, inner_x + (inner_w - ring) / 2, inner_y + (inner_h - ring) / 2, ring, ring,
           fill=None, line=mix(d.primary, d.accent, 0.75), shape=MSO_SHAPE.OVAL, line_width=1.4)
    b.rect(slide, inner_x + (inner_w - ring * 0.62) / 2, inner_y + (inner_h - ring * 0.62) / 2,
           ring * 0.62, ring * 0.62, fill=mix(d.primary, d.accent, 0.28), shape=MSO_SHAPE.OVAL)
    b.rect(slide, inner_x + (inner_w - ring * 0.24) / 2, inner_y + (inner_h - ring * 0.24) / 2,
           ring * 0.24, ring * 0.24, fill=d.accent, shape=MSO_SHAPE.OVAL)

    # corner ticks
    for cx, cy in ((inner_x, inner_y), (inner_x + inner_w - Inches(0.4), inner_y),
                   (inner_x, inner_y + inner_h - Inches(0.4)),
                   (inner_x + inner_w - Inches(0.4), inner_y + inner_h - Inches(0.4))):
        b.rect(slide, cx, cy, Inches(0.4), Emu(9525), fill=mix(d.primary, d.accent, 0.6))
        b.rect(slide, cx, cy, Emu(9525), Inches(0.4), fill=mix(d.primary, d.accent, 0.6))

    if label:
        _, tf = b.textbox(slide, x + inset, y + h - Inches(0.85), w - inset * 2, Inches(0.5))
        on_dark = ensure_contrast(d.paper, d.primary, 6.0)
        b.write(tf, title_case(short(label, 46)), size=11, color=on_dark, bold=True,
                caps=True, tracking=1.6, align=PP_ALIGN.CENTER, first=True)


def layout_image_led(b: DeckBuilder, s: Slide, total: int) -> SlideBox:
    d = b.design
    slide = b.new_slide()
    b.frame(slide, s.kicker, s.title)

    gap = Inches(0.6)
    media_w = Inches(6.1)
    media_x = SLIDE_W - MARGIN - media_w
    text_w = media_x - gap - MARGIN
    top = CONTENT_TOP + Inches(0.25)
    body_h = SLIDE_H - top - Inches(1.25)

    _, tf = b.textbox(slide, MARGIN, top, text_w, Inches(1.6))
    b.write(tf, s.body or s.title, size=d.type_scale["subtitle"], color=d.ink,
            line_spacing=1.35, first=True)

    items = s.bullets[:3]
    y = top + Inches(1.85)
    for i, item in enumerate(items):
        b.rect(slide, MARGIN, y + Inches(0.12), Inches(0.28), Emu(2 * 12700),
               fill=d.accent)
        _, tf = b.textbox(slide, MARGIN + Inches(0.45), y, text_w - Inches(0.45), Inches(1.2))
        b.write(tf, item.head, size=d.type_scale["body"], color=d.ink, font=d.heading_font,
                bold=True, first=True)
        b.write(tf, item.body, size=d.type_scale["small"] + 1, color=d.muted, line_spacing=1.28)
        y += Inches(0.95) + Inches(0.1)
        del i

    _media_panel(b, slide, media_x, top, media_w, body_h, s.media_label)

    b.pattern_band(slide)
    if s.caption:
        _, tf = b.textbox(slide, MARGIN, CONTENT_BOTTOM + Inches(0.06),
                          SLIDE_W - MARGIN * 2, Inches(0.3))
        b.write(tf, s.caption, size=9.5, color=d.muted, italic=True, first=True)
    b.notes(slide.native, s.notes)
    return slide


def layout_summary(b: DeckBuilder, s: Slide, total: int) -> SlideBox:
    d = b.design
    slide = b.new_slide(shade(d.paper, -0.02))
    b.frame(slide, s.kicker, s.title)

    top = CONTENT_TOP + Inches(0.35)
    left_w = Inches(6.4)
    _, tf = b.textbox(slide, MARGIN, top, left_w, Inches(2.2))
    b.write(tf, s.body or s.title, size=d.type_scale["section"], color=d.ink,
            font=d.heading_font, bold=True, line_spacing=1.14, first=True)

    b.rect(slide, MARGIN, top + Inches(2.35), Inches(1.7), Pt(3 * d.rule_weight), fill=d.accent)

    _, tf = b.textbox(slide, MARGIN, top + Inches(2.7), left_w, Inches(0.9))
    b.write(tf, f"{d.brand_name} · {d.preset_label}", size=11,
            color=d.accent if contrast_ratio(d.accent, d.paper) >= 3 else d.ink,
            bold=True, caps=True, tracking=1.6, first=True)
    if d.tagline:
        b.write(tf, title_case(short(d.tagline, 70)), size=12.5, color=d.muted, space_after=0)

    rx = MARGIN + left_w + Inches(0.5)
    rw = SLIDE_W - MARGIN - rx
    items = s.bullets[:3]
    card_h = (SLIDE_H - top - Inches(1.3) - Inches(0.24) * max(0, len(items) - 1)) / max(1, len(items))
    for i, item in enumerate(items):
        y = top + (card_h + Inches(0.24)) * i
        _card(b, slide, rx, y, rw, card_h, fill=d.surface)
        b.rect(slide, rx + Inches(0.32), y + Inches(0.34), Inches(0.46), Inches(0.46),
               fill=d.accent, shape=MSO_SHAPE.OVAL)
        _, tf = b.textbox(slide, rx + Inches(0.32), y + Inches(0.42), Inches(0.46), Inches(0.32))
        b.write(tf, str(i + 1), size=13, color=ensure_contrast("#FFFFFF", d.accent, 4.0),
                font=d.heading_font, bold=True, align=PP_ALIGN.CENTER, first=True)
        _, tf = b.textbox(slide, rx + Inches(1.0), y + Inches(0.32), rw - Inches(1.35),
                          card_h - Inches(0.6))
        b.write(tf, item.head, size=d.type_scale["body"] + 1, color=d.ink, font=d.heading_font,
                bold=True, first=True)
        b.write(tf, item.body, size=d.type_scale["small"] + 1, color=d.muted, line_spacing=1.28)

    b.pattern_band(slide)
    if s.caption:
        _, tf = b.textbox(slide, MARGIN, CONTENT_BOTTOM + Inches(0.06),
                          SLIDE_W - MARGIN * 2, Inches(0.3))
        b.write(tf, s.caption, size=9.5, color=d.muted, italic=True, first=True)
    b.notes(slide.native, s.notes)
    return slide


LAYOUTS = {
    "hero": layout_hero,
    "two_column": layout_two_column,
    "numbered_process": layout_numbered_process,
    "timeline": layout_timeline,
    "metrics": layout_metrics,
    "comparison": layout_comparison,
    "framework": layout_framework,
    "image_led": layout_image_led,
    "quote_led": layout_quote_led,
    "summary": layout_summary,
}


def build_deck(brief, path: Path) -> tuple[int, DesignSystem, list[str]]:
    """Render the full deck to ``path``. Returns (slide count, design, layouts)."""
    design = build_design_system(brief)
    outline = build_outline(brief)
    total = len(outline)

    builder = DeckBuilder(brief)
    builder.total = total
    builder.design = design

    used: list[str] = []
    for slide_spec in outline:
        layout = LAYOUTS.get(slide_spec.layout, layout_two_column)
        box = layout(builder, slide_spec, total)
        builder.index = len(builder.prs.slides)
        used.append(slide_spec.layout)
        del box

    core = builder.prs.core_properties
    core.title = f"{design.brand_name} — {brief.project_name}"
    core.author = "CreateFlow"
    core.subject = short(brief.topic, 180)
    core.keywords = ", ".join([brief.audience[:60], brief.tone[:40], design.preset_label])
    core.comments = (
        f"Generated by CreateFlow from the user brief. Illustrative figures are labelled on-slide. "
        f"Style: {design.preset_label}."
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    builder.prs.save(str(path))
    return total, design, used


def _unused() -> None:  # pragma: no cover - keeps optional imports referenced
    del math, headline_from